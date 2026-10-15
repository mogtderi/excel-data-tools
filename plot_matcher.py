# -*- coding: utf-8 -*-
"""
ابزار تطبیق و تکمیل خودکار اطلاعات قطعات
(Flatten چند ستونی + جستجوی خودکار مساحت و متقاضی)

Coded by Moghtaderi
"""

import os
import threading
import tkinter as tk
import tkinter.font as tkfont
from tkinter import ttk, filedialog, messagebox

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

VILLAGE_COL_NAME = "کد روستا"
PART_COL_PREFIX = "کد قطعه"
AREA_COL_NAME = "مساحت"
APPLICANT_COL_NAME = "متقاضیان"

FONT_NAME = "Arial"
HEADER_COLOR = "305496"
MISSING_COLOR = "FFF2CC"


# ---------------------------------------------------------------------------
# منطق پردازش فایل‌ها (بدون وابستگی به رابط گرافیکی)
# ---------------------------------------------------------------------------
def _norm(v):
    return str(v).strip() if v is not None else ""


def _sheet_with_headers(wb, required_exact, required_prefix=None):
    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue
        headers = [_norm(c.value) for c in ws[1]]
        if all(h in headers for h in required_exact):
            if required_prefix is None or any(h.startswith(required_prefix) for h in headers):
                return ws
    return wb.worksheets[0]


def _build_lookup(ref_ws):
    headers = [_norm(c.value) for c in ref_ws[1]]
    try:
        col_village = headers.index(VILLAGE_COL_NAME)
        col_part = headers.index(PART_COL_PREFIX)
        col_area = headers.index(AREA_COL_NAME)
        col_applicant = headers.index(APPLICANT_COL_NAME)
    except ValueError as e:
        raise ValueError(f"ستون مورد نیاز در فایل مرجع پیدا نشد: {e}")

    lookup = {}
    for row in ref_ws.iter_rows(min_row=2, values_only=True):
        village, part = row[col_village], row[col_part]
        if village is None or part is None:
            continue
        lookup[(_norm(village), _norm(part))] = (row[col_area], row[col_applicant])
    return lookup


def _flatten_target(tgt_ws):
    headers = [_norm(c.value) for c in tgt_ws[1]]
    if VILLAGE_COL_NAME not in headers:
        raise ValueError(f"ستون «{VILLAGE_COL_NAME}» در فایل هدف پیدا نشد.")
    village_idx = headers.index(VILLAGE_COL_NAME)
    part_idxs = [i for i, h in enumerate(headers) if h.startswith(PART_COL_PREFIX)]
    if not part_idxs:
        raise ValueError(f"هیچ ستونی با نام «{PART_COL_PREFIX}» در فایل هدف پیدا نشد.")

    flattened = []
    for row in tgt_ws.iter_rows(min_row=2, values_only=True):
        village = row[village_idx]
        if village is None or _norm(village) == "":
            continue
        village = _norm(village)
        for idx in part_idxs:
            part = row[idx]
            if part is None or _norm(part) == "":
                continue
            flattened.append((village, part))
    return flattened


def _write_output(results, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "نتیجه تکمیل شده"
    ws.sheet_view.rightToLeft = True

    headers = ["کد روستا", "کد قطعه", "مساحت", "متقاضیان"]
    header_fill = PatternFill("solid", start_color=HEADER_COLOR, end_color=HEADER_COLOR)
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
    missing_fill = PatternFill("solid", start_color=MISSING_COLOR, end_color=MISSING_COLOR)
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")

    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, center, border

    for r_idx, (village, part, area, applicant) in enumerate(results, start=2):
        missing = area is None and applicant is None
        vals = [village, part,
                area if area is not None else "یافت نشد",
                applicant if applicant is not None else "یافت نشد"]
        for c_idx, v in enumerate(vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            cell.font = Font(name=FONT_NAME, size=11)
            cell.border = border
            cell.alignment = right if c_idx == 4 else center
            if missing:
                cell.fill = missing_fill

    for i, w in enumerate([14, 12, 14, 42], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    if results:
        ws.auto_filter.ref = f"A1:D{len(results) + 1}"

    wb.save(out_path)


def process_files(ref_path, target_path, out_path, log=lambda msg: None):
    log("در حال خواندن فایل مرجع…")
    ref_wb = openpyxl.load_workbook(ref_path, data_only=True)
    ref_ws = _sheet_with_headers(ref_wb, (VILLAGE_COL_NAME, AREA_COL_NAME, APPLICANT_COL_NAME), PART_COL_PREFIX)
    lookup = _build_lookup(ref_ws)
    log(f"تعداد رکوردهای فایل مرجع: {len(lookup)}")

    log("در حال خواندن و Flatten کردن فایل هدف…")
    tgt_wb = openpyxl.load_workbook(target_path, data_only=True)
    tgt_ws = _sheet_with_headers(tgt_wb, (VILLAGE_COL_NAME,), PART_COL_PREFIX)
    flattened = _flatten_target(tgt_ws)
    log(f"تعداد کدهای قطعه بعد از Flatten: {len(flattened)}")

    results, not_found = [], 0
    for village, part in flattened:
        key = (village, _norm(part))
        area, applicant = lookup.get(key, (None, None))
        if key not in lookup:
            not_found += 1
        results.append((village, part, area, applicant))

    log(f"موارد یافت‌نشده در فایل مرجع: {not_found}")
    log("در حال ساخت فایل خروجی اکسل…")
    _write_output(results, out_path)
    log("پردازش با موفقیت به پایان رسید.")
    return {"total": len(flattened), "not_found": not_found, "matched": len(flattened) - not_found}


# ---------------------------------------------------------------------------
# رابط گرافیکی (Tkinter)
# ---------------------------------------------------------------------------
class App:
    def __init__(self, root):
        self.root = root
        root.title("ابزار تطبیق و تکمیل کد قطعه  —  Coded by Moghtaderi")
        root.geometry("640x600")
        root.resizable(False, False)
        root.configure(bg="#F4F6F9")

        try:
            default_font = tkfont.nametofont("TkDefaultFont")
            default_font.configure(family="Tahoma", size=10)
        except Exception:
            pass

        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass
        style.configure("Accent.TButton", background="#305496", foreground="white",
                         font=("Tahoma", 11, "bold"), padding=8)
        style.map("Accent.TButton", background=[("active", "#25406f")])
        style.configure("TButton", font=("Tahoma", 10), padding=6)
        style.configure("TLabelframe.Label", font=("Tahoma", 10, "bold"), foreground="#305496")

        self.ref_path = tk.StringVar()
        self.target_path = tk.StringVar()

        self._build_header()
        self._build_file_pickers()
        self._build_run_button()
        self._build_log_area()
        self._build_footer()

    # ---- UI sections ----
    def _build_header(self):
        frame = tk.Frame(self.root, bg="#F4F6F9")
        frame.pack(fill="x", padx=20, pady=(18, 6))
        tk.Label(frame, text="ابزار تطبیق و تکمیل خودکار کد قطعه",
                  font=("Tahoma", 15, "bold"), bg="#F4F6F9", fg="#1F2A44").pack(anchor="e")
        tk.Label(frame, text="خواندن فایل چندستونی، Flatten کردن و جستجوی خودکار مساحت / متقاضی",
                  font=("Tahoma", 9), bg="#F4F6F9", fg="#5A6472").pack(anchor="e", pady=(2, 0))

    def _build_file_pickers(self):
        box = ttk.LabelFrame(self.root, text="۱) انتخاب فایل‌ها")
        box.pack(fill="x", padx=20, pady=10)

        self._picker_row(box, "فایل مرجع (خروجی نهایی):", self.ref_path, self._choose_ref)
        self._picker_row(box, "فایل هدف (قطور):", self.target_path, self._choose_target)

    def _picker_row(self, parent, label_text, var, command):
        row = tk.Frame(parent, bg="#FFFFFF")
        row.pack(fill="x", padx=10, pady=8)

        ttk.Button(row, text="انتخاب فایل...", command=command).pack(side="left", padx=(0, 10))
        lbl = tk.Label(row, textvariable=var, font=("Tahoma", 9), fg="#5A6472",
                        bg="#FFFFFF", anchor="e", justify="right")
        lbl.pack(side="left", fill="x", expand=True)
        tk.Label(row, text=label_text, font=("Tahoma", 10, "bold"),
                  bg="#FFFFFF", anchor="e", justify="right", width=22).pack(side="right")
        var.trace_add("write", lambda *a: lbl.config(text=os.path.basename(var.get()) or "فایلی انتخاب نشده"))
        lbl.config(text="فایلی انتخاب نشده")

    def _build_run_button(self):
        frame = tk.Frame(self.root, bg="#F4F6F9")
        frame.pack(fill="x", padx=20, pady=(4, 10))
        self.run_btn = ttk.Button(frame, text="▶  اجرا و ساخت فایل خروجی", style="Accent.TButton",
                                   command=self._run)
        self.run_btn.pack(fill="x")
        self.progress = ttk.Progressbar(frame, mode="indeterminate")
        self.progress.pack(fill="x", pady=(8, 0))

    def _build_log_area(self):
        box = ttk.LabelFrame(self.root, text="۲) گزارش پردازش")
        box.pack(fill="both", expand=True, padx=20, pady=(0, 10))
        self.log_box = tk.Text(box, height=14, font=("Consolas", 9), bg="#0F172A",
                                fg="#D1FAE5", state="disabled", wrap="word")
        self.log_box.pack(fill="both", expand=True, padx=8, pady=8)

    def _build_footer(self):
        tk.Label(self.root, text="Coded by Moghtaderi", font=("Tahoma", 8, "italic"),
                  bg="#F4F6F9", fg="#9AA3AF").pack(pady=(0, 8))

    # ---- File dialogs ----
    def _choose_ref(self):
        path = filedialog.askopenfilename(title="انتخاب فایل مرجع (خروجی نهایی)",
                                           filetypes=[("Excel files", "*.xlsx")])
        if path:
            self.ref_path.set(path)

    def _choose_target(self):
        path = filedialog.askopenfilename(title="انتخاب فایل هدف (قطور)",
                                           filetypes=[("Excel files", "*.xlsx")])
        if path:
            self.target_path.set(path)

    # ---- Run logic ----
    def _run(self):
        if not self.ref_path.get() or not self.target_path.get():
            messagebox.showwarning("هشدار", "لطفاً ابتدا هر دو فایل را انتخاب کنید.")
            return

        default_name = "قطور_تکمیل_شده.xlsx"
        out_path = filedialog.asksaveasfilename(
            title="ذخیره فایل خروجی", defaultextension=".xlsx",
            initialfile=default_name, filetypes=[("Excel files", "*.xlsx")])
        if not out_path:
            return

        self.run_btn.config(state="disabled")
        self.progress.start(10)
        self._clear_log()
        threading.Thread(target=self._run_worker, args=(out_path,), daemon=True).start()

    def _run_worker(self, out_path):
        try:
            stats = process_files(self.ref_path.get(), self.target_path.get(), out_path, log=self._log)
            self.root.after(0, lambda: self._finish(stats, out_path))
        except Exception as e:
            self.root.after(0, lambda: self._error(str(e)))

    # ---- Log helpers (thread-safe via after()) ----
    def _log(self, msg):
        self.root.after(0, lambda: self._append_log(msg))

    def _append_log(self, msg):
        self.log_box.configure(state="normal")
        self.log_box.insert("end", msg + "\n")
        self.log_box.see("end")
        self.log_box.configure(state="disabled")

    def _clear_log(self):
        self.log_box.configure(state="normal")
        self.log_box.delete("1.0", "end")
        self.log_box.configure(state="disabled")

    def _finish(self, stats, out_path):
        self.progress.stop()
        self.run_btn.config(state="normal")
        self._append_log(f"✔ فایل خروجی ذخیره شد: {out_path}")
        messagebox.showinfo(
            "پایان عملیات",
            f"عملیات با موفقیت انجام شد.\n\n"
            f"تعداد کل کدها: {stats['total']}\n"
            f"تطبیق موفق: {stats['matched']}\n"
            f"یافت نشد: {stats['not_found']}"
        )

    def _error(self, msg):
        self.progress.stop()
        self.run_btn.config(state="normal")
        self._append_log(f"✘ خطا: {msg}")
        messagebox.showerror("خطا در پردازش", msg)


if __name__ == "__main__":
    root = tk.Tk()
    App(root)
    root.mainloop()
