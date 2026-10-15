# -*- coding: utf-8 -*-
"""
تشخیص قطعات مشترک بین کارشناسان نقشه‌برداری - نسخه با رابط گرافیکی (GUI)
==========================================================================
همان منطق نسخه‌ی تک‌فایلیِ قبلی، این‌بار با یک پنجره‌ی گرافیکی ساده (Tkinter؛
جزء خود پایتون است و نیاز به نصب کتابخانه‌ی اضافه ندارد) به‌جای خط فرمان.

راهنمای تبدیل به exe در انتهای همین فایل (کامنت) آمده است.
"""

import glob
import os
import platform
import queue
import re
import sys
import threading
import traceback
from datetime import datetime
from itertools import combinations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

VALID_EXT = (".xlsx", ".xls", ".xlsm")

# ==============================================================================
# بخش ۱: نرمال‌سازی متن فارسی  (همان منطق تست‌شده‌ی نسخه‌ی قبلی)
# ==============================================================================
FA_DIGITS = "۰۱۲۳۴۵۶۷۸۹"
AR_DIGITS = "٠١٢٣٤٥٦٧٨٩"


def normalize_fa(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    text = str(value)
    text = text.replace("ي", "ی").replace("ك", "ک")
    text = text.replace("\u200c", " ").replace("\xa0", " ")
    text = re.sub(r"[\u064B-\u065F]", "", text)
    for i, d in enumerate(FA_DIGITS):
        text = text.replace(d, str(i))
    for i, d in enumerate(AR_DIGITS):
        text = text.replace(d, str(i))
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_plot(value):
    text = normalize_fa(value)
    text = text.replace(" ", "")
    text = re.sub(r"\.0$", "", text)
    return text


def normalize_for_match(value):
    text = normalize_fa(value)
    text = re.sub(r"[آأإٱ]", "ا", text)
    text = text.replace(" ", "")
    return text


def expert_from_filename(filepath):
    name = os.path.splitext(os.path.basename(filepath))[0]
    name = re.sub(r"^(اکسل|excel)[_\-\s]*", "", name, flags=re.IGNORECASE)
    name = re.sub(r"[_\-\s]*(اکسل|excel)$", "", name, flags=re.IGNORECASE)
    name = name.strip(" _-")
    return name if name else os.path.basename(filepath)


# ==============================================================================
# بخش ۲: شناسایی ستون‌ها
# ==============================================================================
VILLAGE_KW = ["روستا", "ده "]
COUNTY_KW = ["شهرستان"]
PLOT_KW = ["قطعه", "پلاک"]
EXPERT_KW = ["کارشناس"]
APPLICANT_KW = ["متقاضی", "خانوادگی", "نام"]


def detect_header(df_raw, max_scan_rows=3):
    for row_idx in range(min(max_scan_rows, len(df_raw))):
        row = df_raw.iloc[row_idx]
        mapping = {}
        for col_idx, val in row.items():
            text = normalize_fa(val)
            if not text:
                continue
            if "village" not in mapping and any(kw in text for kw in VILLAGE_KW):
                mapping["village"] = col_idx
                continue
            if "county" not in mapping and any(kw in text for kw in COUNTY_KW):
                mapping["county"] = col_idx
                continue
            if "plot" not in mapping and any(kw in text for kw in PLOT_KW):
                mapping["plot"] = col_idx
                continue
            if "expert" not in mapping and any(kw in text for kw in EXPERT_KW):
                mapping["expert"] = col_idx
                continue
        for col_idx, val in row.items():
            if col_idx in mapping.values():
                continue
            text = normalize_fa(val)
            if not text:
                continue
            if "applicant" not in mapping and any(kw in text for kw in APPLICANT_KW):
                mapping["applicant"] = col_idx
                break
        if "village" in mapping and "plot" in mapping:
            return row_idx, mapping
    return None, None


def guess_columns_by_content(df_raw):
    stats = {}
    for col in df_raw.columns:
        series = df_raw[col].dropna()
        if len(series) == 0:
            continue
        texts = series.astype(str)
        numeric_vals = pd.to_numeric(texts.str.replace(" ", ""), errors="coerce")
        pct_numeric = numeric_vals.notna().mean()
        pct_persian = texts.str.contains(r"[\u0600-\u06FF]").mean()
        uniq_ratio = series.nunique() / len(series)
        stats[col] = dict(pct_numeric=pct_numeric, pct_persian=pct_persian,
                           uniq_ratio=uniq_ratio, avg_val=numeric_vals.mean())

    text_cols = {c: s for c, s in stats.items() if s["pct_persian"] > 0.5}
    village_col = min(text_cols, key=lambda c: text_cols[c]["uniq_ratio"]) if text_cols else None
    remaining_text = {c: s for c, s in text_cols.items() if c != village_col}
    applicant_col = max(remaining_text, key=lambda c: remaining_text[c]["uniq_ratio"]) if remaining_text else None

    numeric_cols = {c: s for c, s in stats.items() if s["pct_numeric"] > 0.5 and s["avg_val"] is not None}
    plot_col = None
    for c, s in sorted(numeric_cols.items(), key=lambda kv: kv[1]["avg_val"]):
        if s["avg_val"] < 100000:
            plot_col = c
            break

    return {"village": village_col, "applicant": applicant_col, "plot": plot_col, "county": None, "expert": None}


def _column_is_mostly_numeric(df_raw, col_idx, data_start):
    series = df_raw[col_idx].iloc[data_start:].dropna()
    if len(series) == 0:
        return False
    numeric_vals = pd.to_numeric(series.astype(str).str.replace(" ", ""), errors="coerce")
    return numeric_vals.notna().mean() > 0.5


def validate_mapping(df_raw, mapping, data_start, filepath, sheet, warnings):
    mapping = dict(mapping)
    for role, label in [("village", "روستا"), ("county", "شهرستان"), ("applicant", "متقاضی"), ("expert", "کارشناس")]:
        col = mapping.get(role)
        if col is None:
            continue
        if _column_is_mostly_numeric(df_raw, col, data_start):
            warnings.append(
                f"در فایل «{os.path.basename(filepath)}» / شیت «{sheet}» ستون «{label}» بیشتر مقادیر عددی "
                f"دارد (احتمالاً با ستون دیگری مثل مبلغ اشتباه گرفته شده)؛ این ستون نادیده گرفته شد."
            )
            mapping[role] = None
    return mapping


# ==============================================================================
# بخش ۳: استخراج رکوردها
# ==============================================================================
def extract_records(filepath, warnings):
    expert_default = expert_from_filename(filepath)
    records = []
    try:
        xls = pd.ExcelFile(filepath)
    except Exception as e:
        warnings.append(f"خطا در باز کردن فایل «{os.path.basename(filepath)}»: {e}")
        return records

    for sheet in xls.sheet_names:
        try:
            df_raw = pd.read_excel(xls, sheet_name=sheet, header=None, dtype=str)
        except Exception as e:
            warnings.append(f"خطا در خواندن شیت «{sheet}» از فایل «{os.path.basename(filepath)}»: {e}")
            continue
        if df_raw.empty:
            continue

        header_row, mapping = detect_header(df_raw)
        if mapping and mapping.get("village") is not None and mapping.get("plot") is not None:
            data_start = header_row + 1
        else:
            mapping = guess_columns_by_content(df_raw)
            data_start = 0
            if mapping.get("village") is None or mapping.get("plot") is None:
                warnings.append(
                    f"در فایل «{os.path.basename(filepath)}» / شیت «{sheet}» ستون روستا یا شماره قطعه "
                    f"شناسایی نشد؛ این شیت نادیده گرفته شد."
                )
                continue

        mapping = validate_mapping(df_raw, mapping, data_start, filepath, sheet, warnings)
        if mapping.get("village") is None or mapping.get("plot") is None:
            warnings.append(
                f"در فایل «{os.path.basename(filepath)}» / شیت «{sheet}» پس از بررسی، ستون روستا یا شماره "
                f"قطعه معتبر نبود؛ این شیت نادیده گرفته شد."
            )
            continue

        for i in range(data_start, len(df_raw)):
            row = df_raw.iloc[i]
            village_raw = row[mapping["village"]] if mapping.get("village") is not None else None
            plot_raw = row[mapping["plot"]] if mapping.get("plot") is not None else None
            village_val = normalize_fa(village_raw)
            plot_val = normalize_plot(plot_raw)
            if not village_val or not plot_val:
                continue

            county_raw = row[mapping["county"]] if mapping.get("county") is not None else None
            applicant_raw = row[mapping["applicant"]] if mapping.get("applicant") is not None else None
            expert_raw = row[mapping["expert"]] if mapping.get("expert") is not None else None
            expert_val = normalize_fa(expert_raw) or expert_default

            records.append({
                "village": village_val,
                "village_match": normalize_for_match(village_raw),
                "county": normalize_fa(county_raw),
                "county_match": normalize_for_match(county_raw),
                "plot": plot_val,
                "applicant": normalize_fa(applicant_raw),
                "expert": expert_val,
                "source_file": os.path.basename(filepath),
                "source_path": os.path.abspath(filepath),
                "sheet": sheet,
                "excel_row": i + 1,
                "has_county": bool(normalize_fa(county_raw)),
            })

    if not records:
        warnings.append(f"هیچ رکورد معتبری در فایل «{os.path.basename(filepath)}» پیدا نشد.")
    return records


# ==============================================================================
# بخش ۴: تشخیص قطعات مشترک
# ==============================================================================
def find_duplicates(df, reference_path=None):
    """
    تشخیص قطعات مشترک.

    حالت عادی (reference_path=None): همان منطق قبلی — هر قطعه‌ای که در بیش
    از یک «کارشناس» (منبع) با مقدار متفاوت دیده شود، تکراری علامت می‌خورد
    (مقایسه‌ی همه‌ی فایل‌ها با همدیگر).

    حالت «فایل مرجع» (reference_path = مسیر کامل یکی از فایل‌های ورودی):
    فقط قطعاتی تکراری در نظر گرفته می‌شوند که هم در فایل مرجع و هم در حداقل
    یکی از فایل‌های دیگر (فرعی) دیده شده باشند. تکرارهای داخلیِ خودِ فایل
    مرجع با خودش، یا تکرار بین دو فایل فرعی بدون حضور فایل مرجع، نادیده
    گرفته می‌شود.
    """
    df = df.copy()
    df["loose_key"] = df["village_match"] + "|" + df["plot"]
    if reference_path:
        df["is_reference"] = df["source_path"] == reference_path

    def is_match(sub):
        if reference_path:
            return sub["is_reference"].any() and (~sub["is_reference"]).any()
        return sub["expert"].nunique() > 1

    groups = []
    for _, grp in df.groupby("loose_key"):
        counties_present = sorted(c for c in grp["county_match"].unique() if c)
        empty_grp = grp[grp["county_match"] == ""]

        if len(counties_present) <= 1:
            if is_match(grp):
                sub = grp.copy()
                sub["confident"] = len(counties_present) == 1 and empty_grp.empty
                groups.append(sub)
        else:
            for county_val in counties_present:
                sub_specific = grp[grp["county_match"] == county_val]
                sub_combined = pd.concat([sub_specific, empty_grp]) if not empty_grp.empty else sub_specific
                if is_match(sub_combined):
                    sub = sub_combined.copy()
                    sub["confident"] = empty_grp.empty
                    groups.append(sub)

    if not groups:
        empty = df.iloc[0:0].copy()
        empty["confident"] = pd.Series(dtype=bool)
        empty["group_id"] = pd.Series(dtype=int)
        return empty

    dup_df = pd.concat(groups, ignore_index=True)
    dup_df["group_key"] = dup_df.apply(
        lambda r: r["loose_key"] + "|" + (r["county_match"] if r["confident"] else "؟"), axis=1
    )
    ordered_keys = sorted(dup_df["group_key"].unique())
    group_ids = {k: i + 1 for i, k in enumerate(ordered_keys)}
    dup_df["group_id"] = dup_df["group_key"].map(group_ids)

    if reference_path:
        # داخل هر گروه، ردیف(های) فایل مرجع اول نمایش داده شود
        dup_df["_ref_order"] = (~dup_df["is_reference"]).astype(int)
        dup_df = dup_df.sort_values(["group_id", "_ref_order", "source_file"]).drop(columns="_ref_order")
    else:
        dup_df = dup_df.sort_values(["group_id", "expert", "source_file"])
    return dup_df.reset_index(drop=True)


def build_pair_summary(dup_df):
    pair_counts = {}
    for _, group in dup_df.groupby("group_id"):
        experts = sorted(group["expert"].unique())
        for e1, e2 in combinations(experts, 2):
            pair_counts[(e1, e2)] = pair_counts.get((e1, e2), 0) + 1
    rows = [
        {"کارشناس ۱": e1, "کارشناس ۲": e2, "تعداد قطعات مشترک": c}
        for (e1, e2), c in sorted(pair_counts.items(), key=lambda kv: -kv[1])
    ]
    return pd.DataFrame(rows)


# ==============================================================================
# بخش ۵: تولید فایل خروجی اکسل
# ==============================================================================
FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F4E78", end_color="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
GROUP_FILL_A = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")
GROUP_FILL_B = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")
UNSURE_FONT = Font(name=FONT_NAME, italic=True, color="C00000", size=10)
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_sheet_rtl(ws):
    ws.sheet_view.rightToLeft = True


def _write_header(ws, headers):
    for c_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24


def write_report(dup_df, pair_df, warnings, out_path, reference_mode=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "قطعات مشترک"
    _style_sheet_rtl(ws)
    headers = ["گروه", "روستا", "شهرستان", "شماره قطعه", "کارشناس", "متقاضی", "نام فایل"]
    if reference_mode:
        headers.append("نوع فایل")
    headers += ["نام شیت", "شماره ردیف در فایل", "توضیح"]
    note_col = len(headers)
    _write_header(ws, headers)

    if dup_df.empty:
        empty_msg = ("هیچ قطعه‌ای از فایل(های) فرعی در فایل مرجع پیدا نشد."
                     if reference_mode else
                     "هیچ قطعه مشترکی بین کارشناسان پیدا نشد.")
        ws.cell(row=2, column=1, value=empty_msg).font = Font(name=FONT_NAME, italic=True)
    else:
        r = 2
        current_group = None
        fill_toggle = True
        for _, row in dup_df.iterrows():
            if row["group_id"] != current_group:
                fill_toggle = not fill_toggle
                current_group = row["group_id"]
            fill = GROUP_FILL_A if fill_toggle else GROUP_FILL_B
            note = "" if row["confident"] else "بدون ستون شهرستان - نیاز به بررسی دستی"
            values = [row["group_id"], row["village"], row["county"] if row["county"] else "—",
                      row["plot"], row["expert"], row["applicant"], row["source_file"]]
            if reference_mode:
                values.append("فایل مرجع" if row["is_reference"] else "فایل فرعی")
            values += [row["sheet"], row["excel_row"], note]
            for c_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=r, column=c_idx, value=val)
                cell.fill = fill
                cell.border = BORDER
                cell.font = UNSURE_FONT if (c_idx == note_col and note) else Font(name=FONT_NAME, size=10.5)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            r += 1

    widths = [8, 18, 16, 14, 20, 26, 22]
    if reference_mode:
        widths.append(14)
    widths += [14, 10, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("خلاصه بین کارشناسان")
    _style_sheet_rtl(ws2)
    _write_header(ws2, ["کارشناس ۱", "کارشناس ۲", "تعداد قطعات مشترک"])
    if pair_df.empty:
        ws2.cell(row=2, column=1, value="موردی یافت نشد.")
    else:
        for r_idx, row in enumerate(pair_df.itertuples(index=False), start=2):
            for c_idx, val in enumerate(row, start=1):
                cell = ws2.cell(row=r_idx, column=c_idx, value=val)
                cell.border = BORDER
                cell.font = Font(name=FONT_NAME, size=10.5)
                cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, w in enumerate([22, 22, 20], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws3 = wb.create_sheet("یادداشت‌های پردازش")
    _style_sheet_rtl(ws3)
    _write_header(ws3, ["ردیف", "پیام"])
    if not warnings:
        ws3.cell(row=2, column=1, value=1)
        ws3.cell(row=2, column=2, value="هیچ هشداری در حین پردازش ثبت نشد.")
    else:
        for idx, msg in enumerate(warnings, start=2):
            ws3.cell(row=idx, column=1, value=idx - 1)
            c = ws3.cell(row=idx, column=2, value=msg)
            c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    ws3.column_dimensions["A"].width = 8
    ws3.column_dimensions["B"].width = 90

    wb.save(out_path)


# ==============================================================================
# بخش ۶: رابط گرافیکی (Tkinter)
# ==============================================================================
class DuplicateFinderApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("ابزار تشخیص قطعات مشترک بین کارشناسان — Coded by Moghtaderi")
        self.geometry("780x720")
        self.minsize(700, 600)

        self.files = []
        self.output_dir = os.path.join(self._base_dir(), "output")
        self.log_queue = queue.Queue()
        self.is_running = False
        self.last_output_path = None
        self.compare_mode = tk.StringVar(value="normal")   # "normal" یا "reference"
        self.reference_choice = tk.StringVar(value="")     # گزینه‌ی انتخاب‌شده در کمبوباکس فایل مرجع

        self._build_style()
        self._build_widgets()
        self.after(100, self._poll_queue)

    def _base_dir(self):
        if getattr(sys, "frozen", False):
            return os.path.dirname(sys.executable)
        return os.path.dirname(os.path.abspath(__file__))

    def _build_style(self):
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except Exception:
            pass
        self.option_add("*Font", ("Tahoma", 10))
        style.configure("Header.TLabel", font=("Tahoma", 13, "bold"))
        style.configure("Accent.TButton", font=("Tahoma", 10, "bold"))

    def _build_widgets(self):
        pad = 10

        ttk.Label(self, text="ابزار تشخیص قطعات مشترک بین کارشناسان", style="Header.TLabel").pack(pady=(pad, 2))
        ttk.Label(self, text="فایل‌های اکسل کارشناسان را اضافه کنید و تحلیل را شروع کنید.").pack(pady=(0, pad))

        # --- فایل‌های ورودی ---
        files_frame = ttk.LabelFrame(self, text="فایل‌های ورودی")
        files_frame.pack(fill="both", expand=False, padx=pad, pady=(0, pad))

        list_container = ttk.Frame(files_frame)
        list_container.pack(fill="both", expand=True, padx=8, pady=8)

        files_scroll = ttk.Scrollbar(list_container, orient="vertical")
        files_scroll.pack(side="right", fill="y")
        self.files_listbox = tk.Listbox(list_container, height=7, selectmode="extended",
                                         yscrollcommand=files_scroll.set)
        self.files_listbox.pack(side="left", fill="both", expand=True)
        files_scroll.config(command=self.files_listbox.yview)

        btns_frame = ttk.Frame(files_frame)
        btns_frame.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Button(btns_frame, text="افزودن فایل‌ها...", command=self.add_files).pack(side="right", padx=4)
        ttk.Button(btns_frame, text="افزودن همه فایل‌های یک پوشه...", command=self.add_folder).pack(side="right", padx=4)
        ttk.Button(btns_frame, text="حذف انتخاب‌شده", command=self.remove_selected).pack(side="right", padx=4)
        ttk.Button(btns_frame, text="پاک‌کردن لیست", command=self.clear_files).pack(side="right", padx=4)

        # --- حالت مقایسه ---
        mode_frame = ttk.LabelFrame(self, text="حالت مقایسه")
        mode_frame.pack(fill="x", padx=pad, pady=(0, pad))
        mode_inner = ttk.Frame(mode_frame)
        mode_inner.pack(fill="x", padx=8, pady=8)

        ttk.Radiobutton(
            mode_inner, text="مقایسه‌ی عادی بین همه‌ی فایل‌ها (هر فایل = یک کارشناس)",
            variable=self.compare_mode, value="normal", command=self._on_mode_change,
        ).pack(anchor="e", fill="x")

        ref_row = ttk.Frame(mode_inner)
        ref_row.pack(anchor="e", fill="x", pady=(6, 0))
        self.reference_combo = ttk.Combobox(
            ref_row, textvariable=self.reference_choice, state="disabled",
            justify="right", width=30,
        )
        self.reference_combo.pack(side="right", padx=(0, 8))
        ttk.Radiobutton(
            ref_row, text="مقایسه‌ی فایل‌های فرعی با یک فایل مرجع (اصلی):",
            variable=self.compare_mode, value="reference", command=self._on_mode_change,
        ).pack(side="right")

        ttk.Label(
            mode_inner,
            text="در این حالت فقط قطعاتی که هم در فایل مرجع و هم در فایل‌های دیگر دیده شوند "
                 "در خروجی می‌آیند؛ تکرارهای داخلیِ خودِ فایل مرجع نادیده گرفته می‌شود.",
            foreground="#666666", font=("Tahoma", 8), justify="right", wraplength=700,
        ).pack(anchor="e", fill="x", pady=(4, 0))

        # --- پوشه خروجی ---
        out_frame = ttk.LabelFrame(self, text="پوشه ذخیره خروجی")
        out_frame.pack(fill="x", padx=pad, pady=(0, pad))
        out_inner = ttk.Frame(out_frame)
        out_inner.pack(fill="x", padx=8, pady=8)
        ttk.Button(out_inner, text="تغییر پوشه...", command=self.choose_output_dir).pack(side="right", padx=4)
        self.output_var = tk.StringVar(value=self.output_dir)
        ttk.Entry(out_inner, textvariable=self.output_var, justify="right").pack(
            side="right", fill="x", expand=True, padx=(4, 4)
        )

        # --- اجرا + نوار پیشرفت ---
        run_frame = ttk.Frame(self)
        run_frame.pack(fill="x", padx=pad, pady=(0, pad))
        self.run_button = ttk.Button(run_frame, text="شروع تحلیل", style="Accent.TButton", command=self.start_analysis)
        self.run_button.pack(side="right", padx=4)
        self.progress = ttk.Progressbar(run_frame, mode="indeterminate")
        self.progress.pack(side="right", fill="x", expand=True, padx=8)

        # --- گزارش پردازش ---
        log_frame = ttk.LabelFrame(self, text="گزارش پردازش")
        log_frame.pack(fill="both", expand=True, padx=pad, pady=(0, pad))
        log_inner = ttk.Frame(log_frame)
        log_inner.pack(fill="both", expand=True, padx=8, pady=8)
        log_scroll = ttk.Scrollbar(log_inner, orient="vertical")
        log_scroll.pack(side="right", fill="y")
        self.log_text = tk.Text(log_inner, height=12, wrap="word", state="disabled",
                                 yscrollcommand=log_scroll.set)
        self.log_text.pack(side="left", fill="both", expand=True)
        log_scroll.config(command=self.log_text.yview)

        # --- دکمه‌های پایانی ---
        bottom_frame = ttk.Frame(self)
        bottom_frame.pack(fill="x", padx=pad, pady=(0, pad))
        self.open_file_btn = ttk.Button(bottom_frame, text="باز کردن فایل خروجی",
                                         command=self.open_output_file, state="disabled")
        self.open_file_btn.pack(side="right", padx=4)
        self.open_folder_btn = ttk.Button(bottom_frame, text="باز کردن پوشه خروجی",
                                           command=self.open_output_folder, state="disabled")
        self.open_folder_btn.pack(side="right", padx=4)

        # --- امضای برنامه ---
        credit_label = ttk.Label(self, text="Coded by Moghtaderi", foreground="#888888",
                                  font=("Tahoma", 8, "italic"))
        credit_label.pack(side="bottom", pady=(0, 6))

    # ---- مدیریت لیست فایل‌ها ----
    def add_files(self):
        paths = filedialog.askopenfilenames(
            title="انتخاب فایل‌های اکسل",
            filetypes=[("Excel files", "*.xlsx *.xls *.xlsm"), ("همه فایل‌ها", "*.*")],
        )
        for p in paths:
            self._add_path(p)

    def add_folder(self):
        folder = filedialog.askdirectory(title="انتخاب پوشه حاوی فایل‌های اکسل")
        if not folder:
            return
        found = [f for f in sorted(glob.glob(os.path.join(folder, "*")))
                 if f.lower().endswith(VALID_EXT) and not os.path.basename(f).startswith("~$")]
        if not found:
            messagebox.showinfo("توجه", "هیچ فایل اکسلی در این پوشه پیدا نشد.")
            return
        for f in found:
            self._add_path(f)

    def _add_path(self, path):
        abspath = os.path.abspath(path)
        if abspath in self.files:
            return
        self.files.append(abspath)
        self.files_listbox.insert("end", os.path.basename(abspath))
        self._refresh_reference_combo()

    def remove_selected(self):
        for idx in reversed(self.files_listbox.curselection()):
            self.files_listbox.delete(idx)
            del self.files[idx]
        self._refresh_reference_combo()

    def clear_files(self):
        self.files_listbox.delete(0, "end")
        self.files = []
        self._refresh_reference_combo()

    def choose_output_dir(self):
        folder = filedialog.askdirectory(title="انتخاب پوشه ذخیره خروجی")
        if folder:
            self.output_var.set(folder)

    # ---- حالت مقایسه / فایل مرجع ----
    def _on_mode_change(self):
        if self.compare_mode.get() == "reference":
            self.reference_combo.config(state="readonly")
            self._refresh_reference_combo()
        else:
            self.reference_combo.config(state="disabled")

    def _refresh_reference_combo(self):
        labels = [f"{i + 1}) {os.path.basename(p)}" for i, p in enumerate(self.files)]
        self.reference_combo["values"] = labels
        if self.reference_choice.get() not in labels:
            self.reference_choice.set(labels[0] if labels else "")

    def _selected_reference_path(self):
        choice = self.reference_choice.get()
        if not choice or ")" not in choice:
            return None
        try:
            idx = int(choice.split(")", 1)[0]) - 1
            return self.files[idx]
        except (ValueError, IndexError):
            return None

    # ---- گزارش زنده ----
    def log(self, msg):
        self.log_queue.put(msg)

    def _poll_queue(self):
        try:
            while True:
                msg = self.log_queue.get_nowait()
                self.log_text.config(state="normal")
                self.log_text.insert("end", msg + "\n")
                self.log_text.see("end")
                self.log_text.config(state="disabled")
        except queue.Empty:
            pass
        self.after(100, self._poll_queue)

    # ---- اجرای تحلیل ----
    def start_analysis(self):
        if self.is_running:
            return
        if not self.files:
            messagebox.showwarning("توجه", "ابتدا حداقل یک فایل اکسل اضافه کنید.")
            return
        out_dir = self.output_var.get().strip()
        if not out_dir:
            messagebox.showwarning("توجه", "پوشه خروجی را مشخص کنید.")
            return

        reference_path = None
        if self.compare_mode.get() == "reference":
            reference_path = self._selected_reference_path()
            if not reference_path:
                messagebox.showwarning("توجه", "یک فایل مرجع از لیست انتخاب کنید.")
                return
            if len(self.files) < 2:
                messagebox.showwarning(
                    "توجه", "برای مقایسه با فایل مرجع، حداقل یک فایل فرعی هم باید اضافه کنید."
                )
                return

        self.is_running = True
        self.run_button.config(state="disabled")
        self.open_file_btn.config(state="disabled")
        self.open_folder_btn.config(state="disabled")
        self.progress.start(12)
        self.log_text.config(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.config(state="disabled")

        threading.Thread(
            target=self._run_worker, args=(list(self.files), out_dir, reference_path), daemon=True
        ).start()

    def _run_worker(self, files, out_dir, reference_path=None):
        try:
            self.log(f"{len(files)} فایل انتخاب شده:")
            for f in files:
                if reference_path:
                    tag = "  [فایل مرجع]" if f == reference_path else "  [فایل فرعی]"
                else:
                    tag = ""
                self.log("  - " + os.path.basename(f) + tag)

            warnings = []
            all_records = []
            for f in files:
                recs = extract_records(f, warnings)
                self.log(f"«{os.path.basename(f)}» -> {len(recs)} رکورد استخراج شد")
                all_records.extend(recs)

            if not all_records:
                self.log("\nهیچ رکورد قابل‌استفاده‌ای در فایل‌ها پیدا نشد.")
                for w in warnings:
                    self.log("هشدار: " + w)
                self._finish(success=False)
                return

            df = pd.DataFrame(all_records)
            dup_df = find_duplicates(df, reference_path=reference_path)
            pair_df = build_pair_summary(dup_df)
            n_groups = dup_df["group_id"].nunique() if not dup_df.empty else 0

            self.log(f"\nتعداد کل رکوردها: {len(df)}")
            if reference_path:
                self.log(
                    f"تعداد قطعاتِ فایل(های) فرعی که در فایل مرجع «{os.path.basename(reference_path)}» هم پیدا شدند: {n_groups}"
                )
            else:
                self.log(f"تعداد قطعات مشترک (گروه): {n_groups}")
            self.log(f"تعداد ردیف‌های مرتبط با قطعات مشترک: {len(dup_df)}")

            os.makedirs(out_dir, exist_ok=True)
            out_path = os.path.join(out_dir, "گزارش_قطعات_مشترک.xlsx")
            if os.path.exists(out_path):
                try:
                    with open(out_path, "a"):
                        pass
                except Exception:
                    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                    out_path = os.path.join(out_dir, f"گزارش_قطعات_مشترک_{stamp}.xlsx")

            write_report(dup_df, pair_df, warnings, out_path, reference_mode=bool(reference_path))
            self.last_output_path = out_path
            self.log(f"\nگزارش نهایی ذخیره شد در:\n{out_path}")
            if warnings:
                self.log(f"\n{len(warnings)} یادداشت پردازشی هم در شیت «یادداشت‌های پردازش» فایل خروجی ثبت شد.")

            self._finish(success=True)
        except Exception:
            self.log("\nخطای غیرمنتظره:\n" + traceback.format_exc())
            self._finish(success=False)

    def _finish(self, success):
        def _update():
            self.progress.stop()
            self.run_button.config(state="normal")
            self.is_running = False
            if success and self.last_output_path:
                self.open_file_btn.config(state="normal")
                self.open_folder_btn.config(state="normal")
                messagebox.showinfo("پایان کار", "تحلیل با موفقیت انجام شد.")
            elif not success:
                messagebox.showerror("خطا", "پردازش با خطا مواجه شد؛ جزئیات را در بخش «گزارش پردازش» ببینید.")
        self.after(0, _update)

    def open_output_file(self):
        if self.last_output_path and os.path.exists(self.last_output_path):
            self._open_path(self.last_output_path)

    def open_output_folder(self):
        if self.last_output_path:
            self._open_path(os.path.dirname(self.last_output_path))

    def _open_path(self, path):
        try:
            if platform.system() == "Windows":
                os.startfile(path)  # noqa
            elif platform.system() == "Darwin":
                os.system(f'open "{path}"')
            else:
                os.system(f'xdg-open "{path}"')
        except Exception as e:
            messagebox.showerror("خطا", f"باز کردن مسیر ممکن نشد:\n{e}")


def main():
    app = DuplicateFinderApp()
    app.mainloop()


if __name__ == "__main__":
    main()

# ==============================================================================
# راهنمای تبدیل این فایل به یک exe ویندوزی با رابط گرافیکی (فقط یک‌بار لازم است):
# ==============================================================================
# ۱) روی یک ویندوز، پایتون ۳ نصب کنید (از python.org) و تیک "Add to PATH" را بزنید.
#    (خود Tkinter جزء پایتون است؛ نیازی به نصب جداگانه ندارد.)
# ۲) در Command Prompt:
#
#      pip install pandas openpyxl xlrd pyinstaller
#      pyinstaller --onefile --windowed --collect-all pandas --collect-all openpyxl ^
#          --collect-all xlrd --name DuplicateFinderGUI duplicate_finder_gui.py
#
#    نکته: --windowed یعنی پنجره سیاه Command Prompt دیگر باز نمی‌شود؛ فقط
#    پنجره‌ی گرافیکی برنامه دیده می‌شود.
#
# ۳) فایل نهایی: dist\DuplicateFinderGUI.exe — همین یک فایل قابل انتقال به هر
#    ویندوز دیگری است، بدون نیاز به نصب پایتون روی آن سیستم.
# ==============================================================================
