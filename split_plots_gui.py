# -*- coding: utf-8 -*-
"""
تفکیک شماره‌قطعه‌های چندتایی به ردیف‌های جداگانه - نسخه با رابط گرافیکی
==========================================================================
همان منطق تست‌شده‌ی split_plots.py، این‌بار با یک پنجره‌ی گرافیکی ساده
(Tkinter؛ جزء خود پایتون است، نیازی به نصب کتابخانه اضافه ندارد).

راهنمای تبدیل به exe در انتهای همین فایل (کامنت) آمده است.
"""

import glob
import os
import platform
import queue
import re
import sys
import threading
import traceback
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

VALID_EXT = (".xlsx", ".xls", ".xlsm")

# ==============================================================================
# بخش ۱: شناسایی ستون «شماره قطعه»  (همان منطق تست‌شده‌ی نسخه‌ی خط‌فرمان)
# ==============================================================================
PLOT_KW = ["قطعه", "پلاک"]
FA_DIGITS = "۰۱۲۳۴۵۶۷۸۹"
AR_DIGITS = "٠١٢٣٤٥٦٧٨٩"
_DIGIT_TRANS = str.maketrans(FA_DIGITS + AR_DIGITS, "01234567890123456789")


def _cell_text(val):
    if val is None:
        return ""
    try:
        if pd.isna(val):
            return ""
    except (TypeError, ValueError):
        pass
    return str(val).replace("\u200c", " ").replace("\xa0", " ").strip()


def detect_plot_column(df_raw, max_scan_rows=3):
    for row_idx in range(min(max_scan_rows, len(df_raw))):
        row = df_raw.iloc[row_idx]
        for col_idx, val in row.items():
            text = _cell_text(val)
            if any(kw in text for kw in PLOT_KW):
                return row_idx, col_idx
    return None, None


def guess_plot_column_by_content(df_raw):
    best_col, best_avg = None, None
    for col in df_raw.columns:
        series = df_raw[col].dropna()
        if len(series) == 0:
            continue
        first_num = series.astype(str).str.translate(_DIGIT_TRANS).str.extract(r"^(\d+)")[0]
        numeric_vals = pd.to_numeric(first_num, errors="coerce")
        if numeric_vals.notna().mean() < 0.5:
            continue
        avg = numeric_vals.mean()
        if avg is None or pd.isna(avg) or avg >= 100000:
            continue
        if best_avg is None or avg < best_avg:
            best_avg, best_col = avg, col
    return best_col


# ==============================================================================
# بخش ۲: شکستن سلول شماره قطعه
# ==============================================================================
def split_plot_cell(raw_value):
    text = _cell_text(raw_value)
    if not text:
        return [raw_value]
    parts = [p.strip() for p in re.split(r"[-+]+", text) if p.strip() != ""]
    if len(parts) <= 1:
        return [raw_value]
    return parts


def looks_like_clean_plot_number(value):
    text = _cell_text(value)
    if not text:
        return True
    normalized = text.translate(_DIGIT_TRANS)
    return bool(re.match(r"^\d+(/\d+)?$", normalized))


# ==============================================================================
# بخش ۳: پردازش یک شیت
# ==============================================================================
def process_sheet(df_raw, sheet_name, source_file, warnings):
    header_row, plot_col = detect_plot_column(df_raw)
    if plot_col is None:
        plot_col = guess_plot_column_by_content(df_raw)
        header_labels = [f"ستون {i + 1}" for i in df_raw.columns]
        if plot_col is None:
            warnings.append(
                f"در فایل «{source_file}» / شیت «{sheet_name}» ستون شماره قطعه شناسایی نشد؛ "
                f"این شیت بدون تغییر کپی شد."
            )
            out_rows = [list(row) + ["", ""] for row in df_raw.itertuples(index=False)]
            return header_labels + ["ردیف اصلی", "توضیح"], out_rows, [False] * len(out_rows)
    else:
        data_start = header_row + 1
        header_labels = [_cell_text(v) or f"ستون {i + 1}" for i, v in df_raw.iloc[header_row].items()]

    if plot_col is not None and header_row is None:
        data_start = 0

    out_rows = []
    is_split_flags = []
    for i in range(data_start, len(df_raw)):
        row = list(df_raw.iloc[i])
        original_row_number = i + 1
        plot_value = row[plot_col]
        parts = split_plot_cell(plot_value)

        if len(parts) == 1:
            note = "" if looks_like_clean_plot_number(parts[0]) else "شماره قطعه نامعمول - نیاز به بررسی دستی"
            out_rows.append(row + [original_row_number, note])
            is_split_flags.append(False)
        else:
            for part in parts:
                new_row = list(row)
                new_row[plot_col] = part
                note = "تفکیک‌شده از ردیف {} (مقدار اصلی: {})".format(original_row_number, _cell_text(plot_value))
                if not looks_like_clean_plot_number(part):
                    note += " - این تکه نامعمول است، نیاز به بررسی دستی"
                out_rows.append(new_row + [original_row_number, note])
                is_split_flags.append(True)

    return header_labels + ["ردیف اصلی", "توضیح"], out_rows, is_split_flags


# ==============================================================================
# بخش ۴: نوشتن فایل خروجی
# ==============================================================================
FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F4E78", end_color="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
SPLIT_FILL = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
NORMAL_FILL = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")
WARN_FONT = Font(name=FONT_NAME, italic=True, color="C00000", size=10)
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_sheet(wb, sheet_name, headers, rows, split_flags, is_first):
    ws = wb.active if is_first else wb.create_sheet()
    ws.title = sheet_name[:31]
    ws.sheet_view.rightToLeft = True

    for c_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24

    note_col_idx = len(headers)
    for r_idx, (row_vals, is_split) in enumerate(zip(rows, split_flags), start=2):
        fill = SPLIT_FILL if is_split else NORMAL_FILL
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.fill = fill
            cell.border = BORDER
            is_note_col = c_idx == note_col_idx
            cell.font = WARN_FONT if (is_note_col and val) else Font(name=FONT_NAME, size=10.5)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for c_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c_idx)].width = 18
    ws.column_dimensions[get_column_letter(note_col_idx)].width = 42


def write_notes_sheet(wb, warnings):
    ws = wb.create_sheet("یادداشت‌های پردازش")
    ws.sheet_view.rightToLeft = True
    for c_idx, title in enumerate(["ردیف", "پیام"], start=1):
        cell = ws.cell(row=1, column=c_idx, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    if not warnings:
        ws.cell(row=2, column=1, value=1)
        ws.cell(row=2, column=2, value="هیچ هشداری در حین پردازش ثبت نشد.")
    else:
        for idx, msg in enumerate(warnings, start=2):
            ws.cell(row=idx, column=1, value=idx - 1)
            c = ws.cell(row=idx, column=2, value=msg)
            c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 90


# ==============================================================================
# بخش ۵: پردازش کامل یک یا چند فایل (برای فراخوانی از GUI)
# ==============================================================================
def run_processing(files, out_path, log):
    warnings = []
    wb = Workbook()
    is_first = [True]

    for filepath in files:
        xls = pd.ExcelFile(filepath)
        base_name = os.path.splitext(os.path.basename(filepath))[0]
        for sheet in xls.sheet_names:
            df_raw = pd.read_excel(xls, sheet_name=sheet, header=None, dtype=object)
            if df_raw.empty:
                continue
            headers, rows, split_flags = process_sheet(df_raw, sheet, os.path.basename(filepath), warnings)
            n_split = sum(split_flags)
            log(f"«{os.path.basename(filepath)}» / شیت «{sheet}»: {len(rows)} ردیف نهایی "
                f"({n_split} ردیف حاصل تفکیک)")
            write_sheet(wb, f"{base_name}-{sheet}"[:31], headers, rows, split_flags, is_first[0])
            is_first[0] = False

    write_notes_sheet(wb, warnings)
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb.save(out_path)
    return warnings


# ==============================================================================
# بخش ۶: رابط گرافیکی (Tkinter)
# ==============================================================================
class SplitPlotsApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("ابزار تفکیک شماره‌قطعه‌های چندتایی — Coded by Moghtaderi")
        self.geometry("780x660")
        self.minsize(700, 560)

        self.files = []
        self.output_dir = os.path.join(self._base_dir(), "output")
        self.log_queue = queue.Queue()
        self.is_running = False
        self.last_output_path = None

        self._build_style()
        self._build_widgets()
        self.after(100, self._poll_queue)

    def _base_dir(self):
        if getattr(sys, "frozen", False):
            return os.path.dirname(sys.executable)
        return os.path.dirname(os.path.abspath(__file__))

    def _build_style(self):
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except Exception:
            pass
        self.option_add("*Font", ("Tahoma", 10))
        style.configure("Header.TLabel", font=("Tahoma", 13, "bold"))
        style.configure("Accent.TButton", font=("Tahoma", 10, "bold"))

    def _build_widgets(self):
        pad = 10

        ttk.Label(self, text="ابزار تفکیک شماره‌قطعه‌های چندتایی", style="Header.TLabel").pack(pady=(pad, 2))
        ttk.Label(
            self, text="فایل‌های اکسل را اضافه کنید؛ سلول‌های چندشماره‌ای (مثل 890-485-172) به ردیف‌های جدا تبدیل می‌شوند."
        ).pack(pady=(0, pad))

        # --- فایل‌های ورودی ---
        files_frame = ttk.LabelFrame(self, text="فایل‌های ورودی")
        files_frame.pack(fill="both", expand=False, padx=pad, pady=(0, pad))

        list_container = ttk.Frame(files_frame)
        list_container.pack(fill="both", expand=True, padx=8, pady=8)

        files_scroll = ttk.Scrollbar(list_container, orient="vertical")
        files_scroll.pack(side="right", fill="y")
        self.files_listbox = tk.Listbox(list_container, height=7, selectmode="extended",
                                         yscrollcommand=files_scroll.set)
        self.files_listbox.pack(side="left", fill="both", expand=True)
        files_scroll.config(command=self.files_listbox.yview)

        btns_frame = ttk.Frame(files_frame)
        btns_frame.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Button(btns_frame, text="افزودن فایل‌ها...", command=self.add_files).pack(side="right", padx=4)
        ttk.Button(btns_frame, text="افزودن همه فایل‌های یک پوشه...", command=self.add_folder).pack(side="right", padx=4)
        ttk.Button(btns_frame, text="حذف انتخاب‌شده", command=self.remove_selected).pack(side="right", padx=4)
        ttk.Button(btns_frame, text="پاک‌کردن لیست", command=self.clear_files).pack(side="right", padx=4)

        # --- خروجی ---
        out_frame = ttk.LabelFrame(self, text="خروجی")
        out_frame.pack(fill="x", padx=pad, pady=(0, pad))
        out_inner = ttk.Frame(out_frame)
        out_inner.pack(fill="x", padx=8, pady=(8, 4))
        ttk.Button(out_inner, text="تغییر پوشه...", command=self.choose_output_dir).pack(side="right", padx=4)
        self.output_var = tk.StringVar(value=self.output_dir)
        ttk.Entry(out_inner, textvariable=self.output_var, justify="right").pack(
            side="right", fill="x", expand=True, padx=(4, 4)
        )
        name_inner = ttk.Frame(out_frame)
        name_inner.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Label(name_inner, text="نام فایل خروجی:").pack(side="right", padx=(4, 4))
        self.filename_var = tk.StringVar(value="تفکیک‌شده.xlsx")
        ttk.Entry(name_inner, textvariable=self.filename_var, justify="right", width=30).pack(side="right")

        # --- اجرا + نوار پیشرفت ---
        run_frame = ttk.Frame(self)
        run_frame.pack(fill="x", padx=pad, pady=(0, pad))
        self.run_button = ttk.Button(run_frame, text="شروع تفکیک", style="Accent.TButton", command=self.start_analysis)
        self.run_button.pack(side="right", padx=4)
        self.progress = ttk.Progressbar(run_frame, mode="indeterminate")
        self.progress.pack(side="right", fill="x", expand=True, padx=8)

        # --- گزارش پردازش ---
        log_frame = ttk.LabelFrame(self, text="گزارش پردازش")
        log_frame.pack(fill="both", expand=True, padx=pad, pady=(0, pad))
        log_inner = ttk.Frame(log_frame)
        log_inner.pack(fill="both", expand=True, padx=8, pady=8)
        log_scroll = ttk.Scrollbar(log_inner, orient="vertical")
        log_scroll.pack(side="right", fill="y")
        self.log_text = tk.Text(log_inner, height=12, wrap="word", state="disabled",
                                 yscrollcommand=log_scroll.set)
        self.log_text.pack(side="left", fill="both", expand=True)
        log_scroll.config(command=self.log_text.yview)

        # --- دکمه‌های پایانی ---
        bottom_frame = ttk.Frame(self)
        bottom_frame.pack(fill="x", padx=pad, pady=(0, pad))
        self.open_file_btn = ttk.Button(bottom_frame, text="باز کردن فایل خروجی",
                                         command=self.open_output_file, state="disabled")
        self.open_file_btn.pack(side="right", padx=4)
        self.open_folder_btn = ttk.Button(bottom_frame, text="باز کردن پوشه خروجی",
                                           command=self.open_output_folder, state="disabled")
        self.open_folder_btn.pack(side="right", padx=4)

        credit_label = ttk.Label(self, text="Coded by Moghtaderi", foreground="#888888",
                                  font=("Tahoma", 8, "italic"))
        credit_label.pack(side="bottom", pady=(0, 6))

    # ---- مدیریت لیست فایل‌ها ----
    def add_files(self):
        paths = filedialog.askopenfilenames(
            title="انتخاب فایل‌های اکسل",
            filetypes=[("Excel files", "*.xlsx *.xls *.xlsm"), ("همه فایل‌ها", "*.*")],
        )
        for p in paths:
            self._add_path(p)

    def add_folder(self):
        folder = filedialog.askdirectory(title="انتخاب پوشه حاوی فایل‌های اکسل")
        if not folder:
            return
        found = [f for f in sorted(glob.glob(os.path.join(folder, "*")))
                 if f.lower().endswith(VALID_EXT) and not os.path.basename(f).startswith("~$")]
        if not found:
            messagebox.showinfo("توجه", "هیچ فایل اکسلی در این پوشه پیدا نشد.")
            return
        for f in found:
            self._add_path(f)

    def _add_path(self, path):
        abspath = os.path.abspath(path)
        if abspath in self.files:
            return
        self.files.append(abspath)
        self.files_listbox.insert("end", os.path.basename(abspath))

    def remove_selected(self):
        for idx in reversed(self.files_listbox.curselection()):
            self.files_listbox.delete(idx)
            del self.files[idx]

    def clear_files(self):
        self.files_listbox.delete(0, "end")
        self.files = []

    def choose_output_dir(self):
        folder = filedialog.askdirectory(title="انتخاب پوشه ذخیره خروجی")
        if folder:
            self.output_var.set(folder)

    # ---- گزارش زنده ----
    def log(self, msg):
        self.log_queue.put(msg)

    def _poll_queue(self):
        try:
            while True:
                msg = self.log_queue.get_nowait()
                self.log_text.config(state="normal")
                self.log_text.insert("end", msg + "\n")
                self.log_text.see("end")
                self.log_text.config(state="disabled")
        except queue.Empty:
            pass
        self.after(100, self._poll_queue)

    # ---- اجرای پردازش ----
    def start_analysis(self):
        if self.is_running:
            return
        if not self.files:
            messagebox.showwarning("توجه", "ابتدا حداقل یک فایل اکسل اضافه کنید.")
            return
        out_dir = self.output_var.get().strip()
        filename = self.filename_var.get().strip()
        if not out_dir:
            messagebox.showwarning("توجه", "پوشه خروجی را مشخص کنید.")
            return
        if not filename:
            messagebox.showwarning("توجه", "نام فایل خروجی را مشخص کنید.")
            return
        if not filename.lower().endswith(".xlsx"):
            filename += ".xlsx"

        self.is_running = True
        self.run_button.config(state="disabled")
        self.open_file_btn.config(state="disabled")
        self.open_folder_btn.config(state="disabled")
        self.progress.start(12)
        self.log_text.config(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.config(state="disabled")

        threading.Thread(
            target=self._run_worker, args=(list(self.files), out_dir, filename), daemon=True
        ).start()

    def _run_worker(self, files, out_dir, filename):
        try:
            self.log(f"{len(files)} فایل انتخاب شده:")
            for f in files:
                self.log("  - " + os.path.basename(f))

            os.makedirs(out_dir, exist_ok=True)
            out_path = os.path.join(out_dir, filename)
            if os.path.exists(out_path):
                try:
                    with open(out_path, "a"):
                        pass
                except Exception:
                    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                    name, ext = os.path.splitext(filename)
                    out_path = os.path.join(out_dir, f"{name}_{stamp}{ext}")

            warnings = run_processing(files, out_path, self.log)

            self.last_output_path = out_path
            self.log(f"\nفایل خروجی ذخیره شد در:\n{out_path}")
            if warnings:
                self.log(f"\n{len(warnings)} یادداشت پردازشی هم در شیت «یادداشت‌های پردازش» فایل خروجی ثبت شد.")

            self._finish(success=True)
        except Exception:
            self.log("\nخطای غیرمنتظره:\n" + traceback.format_exc())
            self._finish(success=False)

    def _finish(self, success):
        def _update():
            self.progress.stop()
            self.run_button.config(state="normal")
            self.is_running = False
            if success and self.last_output_path:
                self.open_file_btn.config(state="normal")
                self.open_folder_btn.config(state="normal")
                messagebox.showinfo("پایان کار", "تفکیک با موفقیت انجام شد.")
            elif not success:
                messagebox.showerror("خطا", "پردازش با خطا مواجه شد؛ جزئیات را در بخش «گزارش پردازش» ببینید.")
        self.after(0, _update)

    def open_output_file(self):
        if self.last_output_path and os.path.exists(self.last_output_path):
            self._open_path(self.last_output_path)

    def open_output_folder(self):
        if self.last_output_path:
            self._open_path(os.path.dirname(self.last_output_path))

    def _open_path(self, path):
        try:
            if platform.system() == "Windows":
                os.startfile(path)  # noqa
            elif platform.system() == "Darwin":
                os.system(f'open "{path}"')
            else:
                os.system(f'xdg-open "{path}"')
        except Exception as e:
            messagebox.showerror("خطا", f"باز کردن مسیر ممکن نشد:\n{e}")


def main():
    app = SplitPlotsApp()
    app.mainloop()


if __name__ == "__main__":
    main()

# ==============================================================================
# راهنمای تبدیل این فایل به یک exe ویندوزی با رابط گرافیکی (فقط یک‌بار لازم است):
# ==============================================================================
# ۱) روی یک ویندوز، پایتون ۳ نصب کنید (از python.org) و تیک "Add to PATH" را بزنید.
# ۲) در Command Prompt:
#
#      pip install pandas openpyxl xlrd pyinstaller
#      pyinstaller --onefile --windowed --collect-all pandas --collect-all openpyxl ^
#          --collect-all xlrd --name SplitPlotsGUI split_plots_gui.py
#
# ۳) فایل نهایی: dist\SplitPlotsGUI.exe
# ==============================================================================
